"""
Resume parser service.

Extracts text from PDF, DOCX, and raw text inputs.
Enforces:
- File size validation (<= max_upload_size_mb).
- Magic bytes and extension validation.
- Safe error handling: corrupted files produce 422, never 500.
"""

from __future__ import annotations

import io
from pathlib import Path

import structlog
from docx import Document
from pypdf import PdfReader

from app.core.config import get_settings
from app.core.exceptions import (
    FileTooLargeError,
    MalformedFileError,
    UnsupportedFileTypeError,
)

log = structlog.get_logger(__name__)

# Magic byte signatures
_PDF_MAGIC = b"%PDF"
_ZIP_MAGIC = b"PK\x03\x04"  # DOCX is a zipped XML package


class ResumeParserService:
    """Service for parsing and validating uploaded resume documents."""

    def __init__(self) -> None:
        self.settings = get_settings()

    def validate_file_size(self, content_length: int) -> None:
        """Reject files exceeding the configured maximum upload limit."""
        if content_length > self.settings.max_upload_size_bytes:
            raise FileTooLargeError(max_mb=self.settings.max_upload_size_mb)

    def detect_file_type(self, filename: str, content: bytes) -> str:
        """
        Validate file extension and magic bytes.

        Returns normalized source_type: 'pdf' or 'docx'.
        Raises UnsupportedFileTypeError if file does not match allowed signatures.
        """
        ext = Path(filename).suffix.lower()

        if ext == ".pdf":
            if not content.startswith(_PDF_MAGIC):
                raise UnsupportedFileTypeError(
                    detail="File has a .pdf extension but invalid PDF magic bytes."
                )
            return "pdf"

        if ext == ".docx":
            if not content.startswith(_ZIP_MAGIC):
                raise UnsupportedFileTypeError(
                    detail="File has a .docx extension but invalid document format."
                )
            return "docx"

        raise UnsupportedFileTypeError(
            detail=f"Unsupported file extension '{ext}'. Only PDF and DOCX files are allowed."
        )

    def parse_pdf(self, content: bytes) -> str:
        """Extract text from PDF file using pypdf."""
        try:
            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise MalformedFileError(detail="Password-protected PDF files are not supported.")

            text_chunks: list[str] = []
            links: list[str] = []
            for _i, page in enumerate(reader.pages):
                page_text = page.extract_text() or ""
                if page_text.strip():
                    text_chunks.append(page_text.strip())

                # Extract hyperlink annotations (e.g. LinkedIn, GitHub, portfolio)
                if "/Annots" in page:
                    for annot in page["/Annots"]:
                        try:
                            obj = annot.get_object()
                            if "/A" in obj and "/URI" in obj["/A"]:
                                uri = str(obj["/A"]["/URI"]).strip()
                                if uri and not uri.startswith("mailto:") and uri not in links:
                                    links.append(uri)
                        except Exception:
                            pass

            if links:
                text_chunks.append("Links:\n" + "\n".join(links))

            extracted = "\n\n".join(text_chunks).strip()
            if not extracted:
                raise MalformedFileError(
                    detail=(
                        "Could not extract readable text from PDF. "
                        "Scanned images are not supported."
                    )
                )
            return extracted
        except (MalformedFileError, UnsupportedFileTypeError):
            raise
        except Exception as exc:
            log.warning("pdf_parsing_failed", error=str(exc))
            raise MalformedFileError(detail=f"Failed to parse PDF document: {exc}") from exc

    def parse_docx(self, content: bytes) -> str:
        """Extract text from DOCX file including paragraphs and tables."""
        try:
            doc = Document(io.BytesIO(content))
            text_chunks: list[str] = []

            # Paragraphs
            for p in doc.paragraphs:
                if p.text.strip():
                    text_chunks.append(p.text.strip())

            # Tables
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        text_chunks.append(" | ".join(row_cells))

            extracted = "\n".join(text_chunks).strip()
            if not extracted:
                raise MalformedFileError(detail="DOCX document appears to be empty.")
            return extracted
        except MalformedFileError:
            raise
        except Exception as exc:
            log.warning("docx_parsing_failed", error=str(exc))
            raise MalformedFileError(detail=f"Failed to parse DOCX document: {exc}") from exc

    def parse_file(self, filename: str, content: bytes) -> tuple[str, str]:
        """
        Validate and extract text from an uploaded resume file.

        Returns tuple of (source_type, extracted_text).
        """
        self.validate_file_size(len(content))
        source_type = self.detect_file_type(filename, content)

        if source_type == "pdf":
            text = self.parse_pdf(content)
        elif source_type == "docx":
            text = self.parse_docx(content)
        else:
            raise UnsupportedFileTypeError()

        return source_type, text

    def clean_text(self, text: str) -> str:
        """Clean and normalize pasted resume text."""
        cleaned = text.strip()
        if len(cleaned) < 30:
            raise MalformedFileError(
                detail="Resume text is too short. Please provide at least 30 characters."
            )
        return cleaned
